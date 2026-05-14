# -*- coding: utf-8 -*-
"""Inspect all DOCX under _articles_import: titles, dates, word counts."""
import json
import os
import re
from datetime import datetime

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMPORT = os.path.join(ROOT, "_articles_import")


def paragraphs_text(doc):
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def guess_title(doc):
    cp = doc.core_properties
    t = (cp.title or "").strip()
    if t:
        return t
    for p in doc.paragraphs[:40]:
        st = (p.style.name if p.style else "") or ""
        if st.startswith("Heading 1") and p.text.strip():
            return p.text.strip()
    for p in doc.paragraphs[:40]:
        if p.text.strip():
            return p.text.strip()[:200]
    return ""


def main():
    rows = []
    for dp, _, fns in os.walk(IMPORT):
        for fn in fns:
            if not fn.lower().endswith(".docx"):
                continue
            path = os.path.join(dp, fn)
            rel = os.path.relpath(path, ROOT)
            try:
                doc = Document(path)
            except Exception as e:
                rows.append({"path": rel, "error": str(e)})
                continue
            title = guess_title(doc)
            cp = doc.core_properties
            created = cp.created
            if created and created.tzinfo:
                created = created.replace(tzinfo=None)
            words = len(re.findall(r"\S+", paragraphs_text(doc)))
            rows.append(
                {
                    "path": rel,
                    "title": title,
                    "created": created.isoformat() if created else None,
                    "words": words,
                    "filename": fn,
                }
            )
    rows.sort(key=lambda r: (r.get("title") or ""))
    out = os.path.join(IMPORT, "_docx_inventory.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print("wrote", out, "count", len(rows))


if __name__ == "__main__":
    main()
