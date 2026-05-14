# -*- coding: utf-8 -*-
"""
Сборка страниц блога из .docx (Pandoc + постобработка).
Источники (по приоритету): articles_docx/ (ручной дроп, UTF-8 имена), затем _articles_import/.
"""
from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pytils.translit import slugify as ru_slugify

ROOT = Path(__file__).resolve().parents[1]
DROP = ROOT / "articles_docx"
FALLBACK = ROOT / "_articles_import"
BLOG = ROOT / "blog"
ASSETS_BLOG = ROOT / "assets" / "blog"
MANIFEST = ROOT / "_articles_import" / "blog_manifest.json"

SITE = "https://lab-itis.ru"
ORG = "ИТиС ЛАБ"

# Сохраняем URL из листинга под ближайшие по смыслу материалы
SLUG_OVERRIDES = {
    "что такое компьютерное зрение на производстве": "kak-rabotaet-mashinnoe-zrenie",
    "когда внедрение компьютерного зрения окупается быстрее всего": "roi-mashinnogo-zreniya",
    "что понимают под системой компьютерного зрения для контроля качества": "defekty-metalla-klassifikatsiya",
}

SKIP_SUBSTR = (
    "заголовок для поисковиков",
    "1. title",
)

AUTHORS = (
    "Дмитрий Новиков",
    "Илья Ковалёв",
    "Елена Петрова",
    "Артём Лебедев",
    "Ольга Морозова",
    "ИТиС ЛАБ",
)


def find_docx_files() -> list[Path]:
    paths: list[Path] = []
    if DROP.is_dir():
        paths.extend(sorted(DROP.rglob("*.docx")))
    if not paths and FALLBACK.is_dir():
        paths.extend(sorted(FALLBACK.rglob("*.docx")))
    # без дубликатов по resolve()
    seen: set[str] = set()
    uniq: list[Path] = []
    for p in paths:
        k = str(p.resolve())
        if k not in seen:
            seen.add(k)
            uniq.append(p)
    return uniq


def word_count(doc: Document) -> int:
    text = "\n".join(p.text for p in doc.paragraphs)
    return len(re.findall(r"\S+", text))


def is_template_doc(doc: Document, path: Path) -> bool:
    if word_count(doc) < 90:
        return True
    t = (doc.core_properties.title or "").lower()
    blob = "\n".join(p.text for p in doc.paragraphs[:8]).lower()
    for s in SKIP_SUBSTR:
        if s in t or s in blob:
            return True
    return False


def _is_separator_line(line: str) -> bool:
    s = line.replace("\u00a0", " ").strip()
    if len(s) < 4:
        return False
    return bool(re.match(r"^[—\-_\s\u2013\u2014]{4,}$", s))


def extract_title_desc(doc: Document) -> tuple[str, str]:
    """Заголовок и description из структуры Word-шаблона или fallback."""
    paras = [p.text.replace("\n", " ").strip() for p in doc.paragraphs]
    title = (doc.core_properties.title or "").strip()
    desc = ""

    def idx(label: str) -> int | None:
        for i, line in enumerate(paras):
            if line.startswith(label):
                return i
        return None

    i = idx("Описание страницы:")
    if i is not None and i + 1 < len(paras) and paras[i + 1]:
        desc = paras[i + 1].strip()

    i = idx("Главный заголовок статьи:")
    if i is None:
        i = idx("Главный заголовок:")
    if i is not None and i + 1 < len(paras) and paras[i + 1]:
        title = paras[i + 1].strip()
    elif title in ("", "Заголовок страницы:"):
        for p in doc.paragraphs[:50]:
            st = (p.style.name if p.style else "") or ""
            if st.startswith("Heading 1") and p.text.strip():
                title = p.text.strip()
                break
        if title in ("", "Заголовок страницы:"):
            for line in paras:
                if (
                    line
                    and not _is_separator_line(line)
                    and line not in (
                        "Заголовок страницы:",
                        "Описание страницы:",
                        "Главный заголовок статьи:",
                        "Главный заголовок:",
                    )
                ):
                    title = line[:200]
                    break

    title = re.sub(r"\s+", " ", title).strip()
    if not desc:
        # первый длинный абзац после заголовка
        hit = False
        for line in paras:
            if line == title:
                hit = True
                continue
            if not hit:
                continue
            if len(line) > 80:
                desc = line
                break
    desc = re.sub(r"\s+", " ", desc).strip()
    if re.match(r"^Главный заголовок\s*:", desc, re.I):
        desc = ""
    if len(desc) > 165:
        desc = desc[:162].rsplit(" ", 1)[0] + "…"
    if len(desc) < 80 and title:
        desc = (title + ". " + desc).strip()[:165]
    if _is_separator_line(title) or len(title.strip()) < 4:
        for line in paras:
            if line and not _is_separator_line(line) and len(line) > 12:
                title = re.sub(r"\s+", " ", line).strip()[:200]
                break
    return title, desc


def make_slug(title: str, used: set[str]) -> str:
    key = title.lower().strip()
    base = SLUG_OVERRIDES.get(key)
    if not base:
        base = ru_slugify(title)
        if not base:
            base = "material"
        if len(base) > 80:
            base = base[:80].rstrip("-")
    if re.match(r"^[\-–—_]+$", base) or len(base.strip("-–—_")) < 2:
        base = "post-" + hashlib.md5(title.encode("utf-8")).hexdigest()[:12]
    s = base
    n = 2
    while s in used:
        s = f"{base}-{n}"
        n += 1
    used.add(s)
    return s


def topic_tag(title: str) -> str:
    t = title.lower()
    if any(x in t for x in ("сиз", "от ", "безопасность", "охрана труда")):
        return "ОТ · СИЗ"
    if "маркиров" in t:
        return "МАРКИРОВКА"
    if any(x in t for x in ("окупа", "невыгодно", "деньг", "потер")):
        return "ЭКОНОМИКА"
    if "трансформац" in t or "цифров" in t:
        return "ЦИФРОВИЗАЦИЯ"
    if "программ" in t and "обеспеч" in t:
        return "SOFTWARE"
    if any(x in t for x in ("зрение", "ии", "камер", "контрол", "качеств", "производств")):
        return "МАШИННОЕ ЗРЕНИЕ"
    return "ЭКСПЕРТИЗА"


def ru_date_short(iso: str) -> str:
    y, m, d = map(int, iso.split("-"))
    months = (
        "",
        "янв",
        "фев",
        "мар",
        "апр",
        "мая",
        "июн",
        "июл",
        "авг",
        "сен",
        "окт",
        "ноя",
        "дек",
    )
    return f"{d} {months[m]} {y}"


def strip_metadata_paragraphs(soup: BeautifulSoup) -> None:
    """Убираем служебные абзацы из альтернативного шаблона (Title/Description/H1)."""
    body = soup.body
    if not body:
        return
    for p in list(body.find_all("p", recursive=False)):
        tx = p.get_text(" ", strip=True)
        if re.match(r"^(Title|Description|H1)\s*:", tx, re.I):
            p.decompose()
            continue
        if tx.startswith("————————"):
            p.decompose()


def demote_h1_to_h2(soup: BeautifulSoup) -> None:
    body = soup.body
    if not body:
        return
    for h1 in list(body.find_all("h1")):
        h2 = soup.new_tag("h2")
        h2.string = h1.get_text(" ", strip=True)
        h1.attrs.pop("id", None)
        h1.replace_with(h2)


def scrub_strong_label_paragraphs(soup: BeautifulSoup) -> None:
    body = soup.body
    if not body:
        return
    labs = {
        "Заголовок страницы",
        "Описание страницы",
        "Главный заголовок статьи",
        "Главный заголовок",
    }
    for p in list(body.find_all("p")):
        tx = p.get_text(" ", strip=True)
        if re.match(r"^[—\-_\s]{6,}$", tx):
            p.decompose()
            continue
        st = p.find("strong")
        if not st:
            continue
        tx0 = st.get_text(" ", strip=True)
        # в Word встречается U+FF1A (полноширинное двоеточие) вместо :
        tx0 = tx0.replace("\uff1a", ":")
        if tx0.startswith("Заголовок страницы") or tx0.startswith("Описание страницы") or tx0.startswith("Главный заголовок"):
            p.decompose()
            continue
        lab = tx0.rstrip(":")
        lab_core = lab.split(":", 1)[0].strip() if ":" in lab else lab
        if lab in labs or lab_core in labs:
            p.decompose()


def scrub_label_h2(soup: BeautifulSoup) -> None:
    for h2 in list(soup.find_all("h2")):
        t = h2.get_text(" ", strip=True)
        for prefix in ("Заголовок страницы:", "Описание страницы:", "Главный заголовок статьи:", "Главный заголовок:"):
            if t.startswith(prefix):
                h2.string = t[len(prefix) :].strip()
                break


def trim_word_noise(soup: BeautifulSoup) -> None:
    """Удаляем служебные абзацы-лейблы (в т.ч. не только первые в body)."""
    body = soup.body
    if not body:
        return
    labels = (
        "Главный заголовок статьи:",
        "Главный заголовок:",
        "Описание страницы:",
        "Заголовок страницы:",
        "Заголовок страницы",
        "Главный заголовок",
    )
    for p in list(body.find_all("p")):
        tx = p.get_text(" ", strip=True)
        if not tx:
            continue
        kill = False
        for lb in labels:
            if tx == lb.rstrip(":") or tx.startswith(lb):
                kill = True
                break
        if kill:
            p.decompose()
    for th in list(body.find_all("th")):
        tx = th.get_text(" ", strip=True).replace(":", "")
        if tx in ("Заголовок страницы", "Описание страницы", "Главный заголовок статьи", "Главный заголовок"):
            tr = th.find_parent("tr")
            if tr:
                tr.decompose()


def promote_strong_sections(soup: BeautifulSoup, page_title: str) -> None:
    body = soup.body
    if not body:
        return
    for p in list(body.find_all("p", recursive=False)):
        st = p.find("strong")
        if not st or len(list(p.children)) != 1:
            continue
        txt = st.get_text(" ", strip=True)
        if not txt or len(txt) > 200:
            continue
        if txt.endswith(":") and len(txt) < 40:
            continue
        if txt == page_title and not p.find_previous_sibling():
            p.decompose()
            continue
        if re.match(r"^\[\s*Фото\s+\d", txt, re.I):
            continue
        h2 = soup.new_tag("h2")
        h2.string = txt
        p.replace_with(h2)


def assign_h2_ids(soup: BeautifulSoup) -> list[tuple[str, str]]:
    body = soup.body
    if not body:
        return []
    toc: list[tuple[str, str]] = []
    n = 0
    for h2 in body.find_all("h2"):
        n += 1
        hid = f"section-{n}"
        h2["id"] = hid
        toc.append((hid, h2.get_text(" ", strip=True)))
    return toc


def photo_placeholders_to_notes(soup: BeautifulSoup) -> None:
    body = soup.body
    if not body:
        return
    for p in list(body.find_all("p", recursive=False)):
        tx = p.get_text(" ", strip=True)
        if re.match(r"^\[\s*Фото\s+\d", tx, re.I):
            div = soup.new_tag("div", attrs={"class": "note"})
            np = soup.new_tag("p")
            np.string = tx
            div.append(np)
            p.replace_with(div)


def run_pandoc(docx: Path, media_root: Path) -> str:
    media_root.mkdir(parents=True, exist_ok=True)
    out = tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        delete=False,
        suffix=".html",
    )
    out.close()
    try:
        subprocess.check_call(
            [
                "pandoc",
                str(docx),
                "-f",
                "docx",
                "-t",
                "html5",
                f"--extract-media={media_root}",
                "-o",
                out.name,
                "--standalone",
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        return Path(out.name).read_text(encoding="utf-8")
    finally:
        Path(out.name).unlink(missing_ok=True)


def relocate_media(
    soup: BeautifulSoup, slug: str, pandoc_media_dir: Path, dest_rel: str
) -> None:
    """pandoc кладёт media в pandoc_media_dir/media/... — копируем в assets/blog/slug/."""
    target_dir = ASSETS_BLOG / slug
    target_dir.mkdir(parents=True, exist_ok=True)
    media_dir = pandoc_media_dir / "media"
    if not media_dir.is_dir():
        return
    files = sorted(media_dir.glob("*"))
    for i, src in enumerate(files, 1):
        if not src.is_file():
            continue
        ext = src.suffix.lower() or ".bin"
        dst_name = f"img-{i:02d}{ext}"
        dst = target_dir / dst_name
        shutil.copy2(src, dst)
    imgs = soup.find_all("img")
    dst_files = sorted(target_dir.glob("img-*"))
    if not dst_files:
        return
    for i, img in enumerate(imgs):
        if i < len(dst_files):
            img["src"] = f"{dest_rel}/{dst_files[i].name}"
        if not img.get("alt"):
            img["alt"] = ""
        img["loading"] = "lazy"
        img["decoding"] = "async"


def article_html_page(
    *,
    title: str,
    description: str,
    slug: str,
    iso_date: str,
    read_min: int,
    tag: str,
    body_inner: str,
    toc: list[tuple[str, str]],
    related: list[dict],
    breadcrumb_short: str,
) -> str:
    url = f"{SITE}/blog/{slug}/"
    og_image = f"{SITE}/assets/images/og-default.jpg"
    headline = title
    if len(headline) < len(title) + 10 and "?" not in title and len(title) < 60:
        headline = title  # ok

    article_ld = {
        "@context": "https://schema.org",
        "@type": "Article",
        "headline": headline,
        "description": description,
        "datePublished": iso_date,
        "author": {"@type": "Organization", "name": ORG},
        "publisher": {"@type": "Organization", "name": ORG, "url": SITE},
        "image": og_image,
        "url": url,
    }
    breadcrumb_ld = {
        "@context": "https://schema.org",
        "@type": "BreadcrumbList",
        "itemListElement": [
            {"@type": "ListItem", "position": 1, "name": "Главная", "item": f"{SITE}/"},
            {"@type": "ListItem", "position": 2, "name": "Блог", "item": f"{SITE}/blog/"},
            {"@type": "ListItem", "position": 3, "name": title[:90], "item": url},
        ],
    }

    toc_html = "\n".join(
        f'              <li><a href="#{hid}">{html.escape(lab)}</a></li>' for hid, lab in toc
    )
    rel_html = ""
    for r in related:
        rt = html.escape(r["title"])
        rtag = html.escape(r["tag"])
        rdate = html.escape(r["date"])
        rel_html += f"""
        <a class="bp" href="/blog/{r["slug"]}/">
          <div class="bp-cover"><div class="bp-cover-placeholder"></div></div>
          <div class="bp-body">
            <div class="bp-meta"><span class="bp-tag">{rtag}</span><span class="bp-date">{rdate}</span></div>
            <div class="bp-title">{rt}</div>
            <div class="bp-read">Читать →</div>
          </div>
        </a>"""

    te = html.escape(title)
    de = html.escape(description)
    be = html.escape(breadcrumb_short)
    return f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <link rel="icon" href="/assets/favicon.svg" type="image/svg+xml">
  <link rel="apple-touch-icon" href="/assets/favicon.svg">
  <title>{te} — {ORG}</title>
  <meta name="description" content="{de}">
  <link rel="canonical" href="{url}">
  <meta property="og:title" content="{te}">
  <meta property="og:description" content="{de}">
  <meta property="og:type" content="article">
  <meta property="og:url" content="{url}">
  <meta property="og:image" content="{og_image}">
  <meta property="article:published_time" content="{iso_date}">
  <meta property="article:author" content="{ORG}">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="{te}">
  <meta name="twitter:description" content="{de}">
  <meta name="twitter:image" content="{og_image}">
  <script type="application/ld+json">
  {json.dumps(article_ld, ensure_ascii=False, indent=2)}
  </script>
  <script type="application/ld+json">
  {json.dumps(breadcrumb_ld, ensure_ascii=False, indent=2)}
  </script>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link rel="preload" as="style" href="https://fonts.googleapis.com/css2?family=Unbounded:wght@700;900&family=Manrope:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap">
  <link href="https://fonts.googleapis.com/css2?family=Unbounded:wght@700;900&family=Manrope:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
  <link rel="stylesheet" href="/css/style.css">
</head>
<body>

<div id="nav-root"></div>
<div id="modals-root"></div>

<div class="page-top">

  <div class="article-hero">
    <div class="wrap">
      <div class="breadcrumb">
        <a href="/">Главная</a>
        <span class="sep">/</span>
        <a href="/blog/">Блог</a>
        <span class="sep">/</span>
        <span class="cur">{be}</span>
      </div>
      <div class="article-meta">
        <span class="article-tag">{html.escape(tag)}</span>
        <span class="article-date">{ru_date_short(iso_date)}</span>
        <span class="article-read">{read_min} мин. чтения</span>
      </div>
      <h1 class="rv">{te}</h1>
    </div>
  </div>

  <div class="wrap">
    <div class="article-cover">
      <div class="article-cover-placeholder" role="img" aria-label="Обложка статьи"></div>
    </div>
  </div>

  <section class="section" style="border:none;padding-top:0;">
    <div class="wrap">
      <div class="article-body">

        <div class="article-prose rv">
{body_inner}
          <div class="article-cta">
            <div class="cta-strip" style="border:none;padding:0;">
              <h2 style="font-size:clamp(22px,2.6vw,36px);">Нужна такая система на вашем производстве?</h2>
              <button class="btn btn-w" onclick="openForm('simple')">Обсудить проект</button>
            </div>
          </div>
        </div>

        <aside class="article-sidebar rv d2">
          <div class="sidebar-toc">
            <div class="toc-title">Содержание</div>
            <ul class="toc-list">
{toc_html}
            </ul>
          </div>
        </aside>

      </div>
    </div>
  </section>

  <div class="related-articles">
    <div class="wrap">
      <div class="sec-hd rv">
        <div>
          <div class="mono sec-num">Ещё материалы</div>
          <h2>Читайте также</h2>
        </div>
        <a href="/blog/" class="btn btn-g">Все статьи →</a>
      </div>
      <div class="blog-grid rv">
{rel_html}
      </div>
    </div>
  </div>

</div>

<div id="footer-root"></div>

<script src="/js/data.js"></script>
<script src="/js/components.js"></script>
<script src="/js/main.js"></script>
<script>
  initComponents();
  initPage();
</script>
</body>
</html>
"""


def stable_date(title: str, idx: int) -> str:
    base = date(2025, 2, 1)
    h = int(hashlib.md5(title.encode("utf-8")).hexdigest()[:6], 16)
    d = base.toordinal() + (h % 85) + idx
    return date.fromordinal(d).isoformat()


def pick_related(all_posts: list[dict], slug: str, n: int = 3) -> list[dict]:
    others = [p for p in all_posts if p["slug"] != slug]
    if not others:
        return []
    h = int(hashlib.md5(slug.encode()).hexdigest(), 16)
    out = []
    for j in range(min(n, len(others))):
        out.append(others[(h + j) % len(others)])
    return out


def process_one(
    docx: Path,
    *,
    used_slugs: set[str],
    iso_date: str,
    author: str,
) -> dict | None:
    doc = Document(str(docx))
    if is_template_doc(doc, docx):
        return None
    title, desc = extract_title_desc(doc)
    if not title or len(title) < 8:
        return None
    wc = word_count(doc)
    read_min = max(4, round(wc / 220))
    slug = make_slug(title, used_slugs)
    tag = topic_tag(title)
    tmp_media = ROOT / "_articles_import" / "_pandoc_tmp" / slug
    if tmp_media.exists():
        shutil.rmtree(tmp_media)
    html_raw = run_pandoc(docx, tmp_media)
    soup = BeautifulSoup(html_raw, "lxml")
    strip_metadata_paragraphs(soup)
    demote_h1_to_h2(soup)
    promote_strong_sections(soup, title)
    photo_placeholders_to_notes(soup)
    trim_word_noise(soup)
    scrub_strong_label_paragraphs(soup)
    scrub_label_h2(soup)
    toc = assign_h2_ids(soup)
    body = soup.body
    if not body:
        return None
    dest_rel = f"/assets/blog/{slug}"
    relocate_media(soup, slug, tmp_media, dest_rel)
    inner = body.decode_contents()
    if tmp_media.exists():
        shutil.rmtree(tmp_media, ignore_errors=True)

    bc = title if len(title) <= 48 else title[:45] + "…"
    return {
        "slug": slug,
        "title": title,
        "description": desc,
        "iso_date": iso_date,
        "read_min": read_min,
        "tag": tag,
        "author": author,
        "body_inner": inner,
        "toc": toc,
        "breadcrumb_short": bc,
        "words": wc,
    }


def clean_listing_blurb(text: str, title: str) -> str:
    """Убираем служебные префиксы из лидов карточек листинга."""
    s = re.sub(r"\s+", " ", (text or "").strip())
    for prefix in ("Title:", "title:", "Description:", "description:", "H1:", "h1:"):
        if s.lower().startswith(prefix.lower()):
            s = s[len(prefix) :].strip()
    if re.match(r"^[—\-_\s\u2013\u2014.]{4,}$", s):
        s = ""
    parts = re.split(r"\s*\.\s*", s)
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) >= 2 and len(set(parts)) == 1:
        s = parts[0] + "."
    if len(s) < 40 and title:
        s = (title[:120] + ("" if len(title) <= 120 else "…")).strip()
    return s


def write_listing(posts: list[dict]) -> None:
    posts = sorted(posts, key=lambda p: p["iso_date"], reverse=True)
    n = len(posts)
    cards = []
    for i, p in enumerate(posts):
        tag_u = html.escape(p["tag"].upper())
        d = html.escape(ru_date_short(p["iso_date"]))
        tt = html.escape(p["title"])
        blur = clean_listing_blurb(p.get("description") or "", p.get("title") or "")
        td = html.escape(blur)
        cards.append(
            f"""
        <a class="bp-a" href="/blog/{p["slug"]}/" data-topic="{p["slug"][:12]}">
          <div class="bp-a-c">
            <div class="bp-a-c-g"></div>
            <div class="bp-a-tag">{tag_u}</div>
          </div>
          <div class="bp-a-body">
            <div class="bp-a-meta"><span>{d}</span><span>{p["read_min"]} мин</span></div>
            <div class="bp-a-t">{tt}</div>
            <div class="bp-a-d">{td}</div>
          </div>
        </a>"""
        )
    listing = f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <link rel="icon" href="/assets/favicon.svg" type="image/svg+xml">
  <link rel="apple-touch-icon" href="/assets/favicon.svg">
  <title>Блог — {ORG}</title>
  <meta name="description" content="Статьи о машинном зрении, промышленной автоматизации и ИИ для производства от команды {ORG}.">
  <link rel="canonical" href="{SITE}/blog/">
  <meta property="og:title" content="Блог — {ORG}">
  <meta property="og:description" content="Статьи о машинном зрении, промышленной автоматизации и ИИ для производства.">
  <meta property="og:type" content="website">
  <meta property="og:url" content="{SITE}/blog/">
  <meta property="og:image" content="{SITE}/assets/images/og-default.jpg">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="Блог — {ORG}">
  <meta name="twitter:description" content="Статьи о машинном зрении, промышленной автоматизации и ИИ для производства.">
  <meta name="twitter:image" content="{SITE}/assets/images/og-default.jpg">
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link rel="preload" as="style" href="https://fonts.googleapis.com/css2?family=Unbounded:wght@700;900&family=Manrope:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap">
  <link href="https://fonts.googleapis.com/css2?family=Unbounded:wght@700;900&family=Manrope:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
  <link rel="stylesheet" href="/css/style.css">
    <style>
    .bg-a {{ display: grid; grid-template-columns: repeat(3,1fr); gap: 2px; align-items: stretch; }}
    @media(max-width:1000px) {{ .bg-a {{ grid-template-columns: repeat(2,1fr); }} }}
    @media(max-width:600px) {{ .bg-a {{ grid-template-columns: 1fr; }} }}
    .bp-a {{ height: 100%; min-height: 0; background: var(--s1); display: flex; flex-direction: column; justify-content: flex-start; cursor: pointer; transition: background .2s; text-decoration: none; color: inherit; box-sizing: border-box; }}
    .bp-a:hover {{ background: var(--s2); }}
    .bp-a-c {{ flex: 0 0 auto; aspect-ratio: 16/10; background: var(--s2); position: relative; overflow: hidden; }}
    .bp-a-c-g {{ position: absolute; inset: 0; background-image: linear-gradient(rgba(255,255,255,.025) 1px,transparent 1px), linear-gradient(90deg,rgba(255,255,255,.025) 1px,transparent 1px); background-size: 32px 32px; transition: filter .25s; }}
    .bp-a:hover .bp-a-c-g {{ filter: brightness(1.4); }}
    .bp-a-tag {{ position: absolute; top: 14px; left: 14px; font-family: var(--fm); font-size: 10px; color: var(--ok); letter-spacing: .1em; background: rgba(9,9,9,.7); backdrop-filter: blur(8px); padding: 6px 10px; border: 1px solid var(--line2); }}
    .bp-a-body {{ flex: 1 1 auto; min-height: 0; padding: 24px 24px 28px; display: flex; flex-direction: column; gap: 12px; justify-content: flex-start; }}
    .bp-a-meta {{ flex: 0 0 auto; font-family: var(--fm); font-size: 10.5px; color: var(--m); letter-spacing: .06em; display: flex; justify-content: space-between; }}
    .bp-a-t {{ flex: 0 0 auto; font-family: var(--fh); font-weight: 700; font-size: 18px; line-height: 1.2; letter-spacing: -.01em; display: -webkit-box; -webkit-line-clamp: 3; -webkit-box-orient: vertical; overflow: hidden; min-height: 3.6em; }}
    .bp-a-d {{ flex: 0 0 auto; min-height: calc(1.55em * 3); color: var(--m2); font-size: 14px; line-height: 1.55; display: -webkit-box; -webkit-line-clamp: 3; -webkit-box-orient: vertical; overflow: hidden; }}
  </style>
</head>
<body>

<div id="nav-root"></div>
<div id="modals-root"></div>

<div class="page-top">

  <section class="section" id="blog">
    <div class="wrap">
      <div class="section-num">01 / СТАТЬИ</div>
      <h2 class="section-h" style="margin-bottom:48px;margin-top:16px;">{n} материалов</h2>

      <div class="bg-a" id="blog-grid">
{"".join(cards)}
      </div>
    </div>
  </section>

  <section class="section" style="padding:64px 0;border-bottom:none;">
    <div class="wrap">
      <div class="cta-strip rv">
        <h2>Есть вопрос по вашему производству?</h2>
        <button class="btn btn-w" onclick="openForm('simple')">Проконсультироваться</button>
      </div>
    </div>
  </section>

</div>

<div id="footer-root"></div>

<script src="/js/data.js"></script>
<script src="/js/components.js"></script>
<script src="/js/main.js"></script>
<script>
  initComponents();
  initPage();
</script>
</body>
</html>
"""
    BLOG.joinpath("index.html").write_text(listing, encoding="utf-8")


def main() -> None:
    ASSETS_BLOG.mkdir(parents=True, exist_ok=True)
    files = find_docx_files()
    if not files:
        raise SystemExit("Нет .docx: положите файлы в articles_docx/ или _articles_import/")

    # дедуп по нормализованному заголовку (две «Цифровая трансформация…»)
    seen_titles: set[str] = set()
    queue: list[Path] = []
    for p in files:
        doc = Document(str(p))
        if is_template_doc(doc, p):
            continue
        title, _ = extract_title_desc(doc)
        key = title.lower().strip()
        if key in seen_titles:
            continue
        seen_titles.add(key)
        queue.append(p)

    used_slugs: set[str] = set()
    built: list[dict] = []
    for idx, docx in enumerate(queue):
        iso = stable_date(docx.name, idx)
        author = AUTHORS[idx % len(AUTHORS)]
        rec = process_one(docx, used_slugs=used_slugs, iso_date=iso, author=author)
        if not rec:
            continue
        built.append(rec)

    # второй проход: related + запись файлов
    for rec in built:
        rel = pick_related(
            [{"slug": x["slug"], "title": x["title"], "tag": x["tag"], "date": ru_date_short(x["iso_date"])} for x in built],
            rec["slug"],
            3,
        )
        page = article_html_page(
            title=rec["title"],
            description=rec["description"],
            slug=rec["slug"],
            iso_date=rec["iso_date"],
            read_min=rec["read_min"],
            tag=rec["tag"],
            body_inner="\n" + rec["body_inner"] + "\n",
            toc=rec["toc"],
            related=rel,
            breadcrumb_short=rec["breadcrumb_short"],
        )
        out_dir = BLOG / rec["slug"]
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / "index.html").write_text(page, encoding="utf-8")

    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(built, ensure_ascii=False, indent=2), encoding="utf-8")
    write_listing(built)
    print("posts", len(built), "manifest", MANIFEST)


if __name__ == "__main__":
    main()
